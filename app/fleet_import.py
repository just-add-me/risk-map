"""Разбор файлов со списком судов: CSV / TXT / XLSX.

Понимает заголовки (mmsi, imo, name/vessel/ship) на русском и английском, а
если заголовков нет — угадывает по длине чисел: 9 цифр — MMSI, 7 — IMO.
"""
from __future__ import annotations

import csv
import io
import re
from typing import Iterable

HEADER_ALIASES = {
    "mmsi": {"mmsi", "ммси"},
    "imo": {"imo", "имо", "imo number", "imo no", "imo_no", "imo#"},
    "name": {"name", "vessel", "vessel name", "ship", "ship name", "название", "судно", "имя"},
}


def _norm_header(h: str) -> str:
    h = (h or "").strip().lower()
    for key, aliases in HEADER_ALIASES.items():
        if h in aliases:
            return key
    return ""


def clean_mmsi(v) -> str:
    s = re.sub(r"\D", "", str(v or "").split(".")[0])
    return s if len(s) == 9 else ""


def clean_imo(v) -> str:
    s = re.sub(r"\D", "", str(v or "").split(".")[0])
    return s if len(s) == 7 else ""


def imo_valid(imo: str) -> bool:
    """У номера IMO есть контрольная цифра: сумма первых шести цифр, умноженных на 7…2, по модулю 10."""
    if not re.fullmatch(r"\d{7}", imo or ""):
        return False
    return sum(int(d) * w for d, w in zip(imo[:6], range(7, 1, -1))) % 10 == int(imo[6])


def _rows_from_table(rows: Iterable[list]) -> list[dict]:
    rows = [[("" if c is None else str(c)).strip() for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return []
    header = [_norm_header(c) for c in rows[0]]
    out = []
    if any(header):
        for r in rows[1:]:
            item = {k: r[i] for i, k in enumerate(header) if k and i < len(r)}
            out.append(item)
        return out
    # без заголовков: угадываем по содержимому
    for r in rows:
        item: dict = {}
        words: list[str] = []
        for c in r:
            for tok in c.split():
                bare = re.sub(r"^(mmsi|imo)[:#]?", "", tok, flags=re.I)
                if re.fullmatch(r"\d{9}", bare) and not item.get("mmsi"):
                    item["mmsi"] = bare
                elif re.fullmatch(r"\d{7}", bare) and not item.get("imo"):
                    item["imo"] = bare
                elif not re.fullmatch(r"[\d.]+", tok):
                    words.append(tok)
        if words:
            item["name"] = " ".join(words)
        out.append(item)
    return out


def parse_file(data: bytes, filename: str) -> list[dict]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = list(wb.worksheets[0].iter_rows(values_only=True))
        return _normalize(_rows_from_table([list(r) for r in rows]))
    if name.endswith(".xls"):
        raise ValueError("Старый формат .xls не поддерживается — сохраните файл как .xlsx или .csv")
    text = data.decode("utf-8-sig", errors="ignore")
    try:
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t|")
        rows = list(csv.reader(io.StringIO(text), dialect))
    except csv.Error:
        rows = [re.split(r"[,;\t|]+", ln) for ln in text.splitlines()]
    return _normalize(_rows_from_table(rows))


def _normalize(items: list[dict]) -> list[dict]:
    out = []
    for it in items:
        out.append({
            "mmsi": clean_mmsi(it.get("mmsi")),
            "imo": clean_imo(it.get("imo")),
            "name": (it.get("name") or "").strip()[:80],
        })
    return [o for o in out if o["mmsi"] or o["imo"] or o["name"]]
